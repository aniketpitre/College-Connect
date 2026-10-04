"""
Accounts reports (spec R4, §3.4): day book, head-wise and mode-wise collection, outstanding
dues by class, defaulters and scholarship receivables, each also as an Excel file.

Collection reports count receipts by the day they were issued (India time). A receipt that was
later cancelled is listed in the day book as cancelled and left out of the collection totals.
"""

import io
from datetime import date
from typing import Any

from bson import ObjectId
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core.db import get_db
from app.core.errors import AppError
from app.modules.fees import ledger
from app.modules.fees.receipts import IST, MODES, ist_day_range
from app.modules.fees.service import oid
from app.modules.students import service as students

# Column kinds: "text", "money" (paise → rupees in Excel), "int", "date".
Column = tuple[str, str, str]  # key, header, kind


def _range(date_from: str | None, date_to: str | None) -> dict[str, Any]:
    start = ist_day_range(date_from)[0]
    end = ist_day_range(date_to or date_from)[1]
    if end <= start:
        raise AppError(422, "The end date is before the start date.", field="date_to")
    return {"$gte": start, "$lt": end}


def day_book(date_from: str | None, date_to: str | None) -> dict[str, Any]:
    rows = list(get_db().receipts.find({"collected_at": _range(date_from, date_to)}).sort("collected_at", 1))
    valid = [r for r in rows if r["status"] == "valid"]
    by_mode: dict[str, int] = {}
    for r in valid:
        by_mode[MODES[r["mode"]]] = by_mode.get(MODES[r["mode"]], 0) + r["amount"]
    return {
        "title": "Day book",
        "columns": [
            ("time", "Date & time", "text"),
            ("number", "Receipt", "text"),
            ("prn", "PRN", "text"),
            ("student", "Student", "text"),
            ("class", "Class", "text"),
            ("mode", "Mode", "text"),
            ("reference", "Reference", "text"),
            ("collector", "Collected by", "text"),
            ("status", "Status", "text"),
            ("amount", "Amount", "money"),
        ],
        "rows": [
            {
                "time": r["collected_at"].astimezone(IST).strftime("%d-%m-%Y %H:%M"),
                "number": r["number"],
                "prn": r["student"]["prn"],
                "student": r["student"]["name"],
                "class": r["student"]["class"],
                "mode": MODES[r["mode"]],
                "reference": r.get("reference", ""),
                "collector": r["collector_name"],
                "status": "Cancelled" if r["status"] == "cancelled" else "",
                "amount": r["amount"] if r["status"] == "valid" else 0,
            }
            for r in rows
        ],
        "totals": {
            "Collected": sum(r["amount"] for r in valid),
            **by_mode,
            "Cancelled receipts": len(rows) - len(valid),
        },
    }


def head_wise(date_from: str | None, date_to: str | None) -> dict[str, Any]:
    totals: dict[str, int] = {}
    names: dict[str, str] = {}
    for r in get_db().receipts.find({"collected_at": _range(date_from, date_to), "status": "valid"}):
        for ln in r["lines"]:
            totals[ln["code"]] = totals.get(ln["code"], 0) + ln["amount"]
            names[ln["code"]] = ln["name"]
    rows = [{"code": c, "name": names[c], "amount": a} for c, a in sorted(totals.items(), key=lambda x: -x[1])]
    return {
        "title": "Head-wise collection",
        "columns": [("code", "Head", "text"), ("name", "Name", "text"), ("amount", "Collected", "money")],
        "rows": rows,
        "totals": {"Total": sum(totals.values())},
    }


def mode_wise(date_from: str | None, date_to: str | None) -> dict[str, Any]:
    agg: dict[str, dict[str, int]] = {}
    for r in get_db().receipts.find({"collected_at": _range(date_from, date_to), "status": "valid"}):
        row = agg.setdefault(r["mode"], {"count": 0, "amount": 0})
        row["count"] += 1
        row["amount"] += r["amount"]
    rows = [
        {"mode": MODES[m], "count": v["count"], "amount": v["amount"]}
        for m, v in sorted(agg.items(), key=lambda x: -x[1]["amount"])
    ]
    return {
        "title": "Mode-wise collection",
        "columns": [("mode", "Mode", "text"), ("count", "Receipts", "int"), ("amount", "Collected", "money")],
        "rows": rows,
        "totals": {"Total": sum(r["amount"] for r in rows)},
    }


def _year_accounts(year_id: ObjectId, student_filter: dict[str, Any]) -> list[tuple[dict[str, Any], dict[str, Any]]]:
    """(student, ledger summary) for every student with entries in the year, matching the filter."""
    db = get_db()
    by_student: dict[ObjectId, list[dict[str, Any]]] = {}
    for e in db.ledger_entries.find({"academic_year_id": year_id}).sort("at", 1):
        by_student.setdefault(e["student_id"], []).append(e)
    roster = db.students.find({"_id": {"$in": list(by_student)}, **student_filter}).sort("prn", 1)
    return [(s, ledger.summary(by_student[s["_id"]])) for s in roster]


def _class_filter(programme_id: str | None, year_of_study: int | None) -> dict[str, Any]:
    f: dict[str, Any] = {}
    if programme_id:
        f["programme_id"] = oid(programme_id, "Programme")
    if year_of_study:
        f["year_of_study"] = year_of_study
    return f


def outstanding(year_id: ObjectId, programme_id: str | None, year_of_study: int | None) -> dict[str, Any]:
    lookups = students._lookups()
    rows = []
    for s, acc in _year_accounts(year_id, _class_filter(programme_id, year_of_study)):
        sm = students.summary(s, lookups)
        rows.append(
            {
                "prn": s["prn"],
                "student": s["name"],
                "class": " ".join(x for x in (sm["programme_code"], sm["year_label"], sm["division"]) if x),
                "category": sm["category_code"] or "",
                "fee": acc["demand"] + acc["charges"],
                "paid": acc["paid"],
                "relief": acc["concessions"] + acc["scholarships"],
                "balance": acc["balance"],
            }
        )
    return {
        "title": "Outstanding dues",
        "columns": [
            ("prn", "PRN", "text"),
            ("student", "Student", "text"),
            ("class", "Class", "text"),
            ("category", "Category", "text"),
            ("fee", "Fee + charges", "money"),
            ("paid", "Paid", "money"),
            ("relief", "Concessions + scholarships", "money"),
            ("balance", "Balance", "money"),
        ],
        "rows": rows,
        "totals": {
            "Students": len(rows),
            "Fee + charges": sum(r["fee"] for r in rows),
            "Paid": sum(r["paid"] for r in rows),
            "Balance": sum(r["balance"] for r in rows),
        },
    }


def defaulters(
    year_id: ObjectId, programme_id: str | None, year_of_study: int | None, as_of: str | None
) -> dict[str, Any]:
    on = date.fromisoformat(as_of) if as_of else date.today()
    lookups = students._lookups()
    db = get_db()
    by_student: dict[ObjectId, list[dict[str, Any]]] = {}
    for e in db.ledger_entries.find({"academic_year_id": year_id}).sort("at", 1):
        by_student.setdefault(e["student_id"], []).append(e)
    rows = []
    for s in db.students.find({"_id": {"$in": list(by_student)}, **_class_filter(programme_id, year_of_study)}).sort(
        "prn", 1
    ):
        acc = ledger.summary(by_student[s["_id"]], today=on)
        if acc["overdue"] <= 0:
            continue
        sm = students.summary(s, lookups)
        oldest = next((i for i in acc["installments"] if i["overdue"]), None)
        rows.append(
            {
                "prn": s["prn"],
                "student": s["name"],
                "class": " ".join(x for x in (sm["programme_code"], sm["year_label"], sm["division"]) if x),
                "phone": s.get("phone") or "",
                "guardian_phone": (s.get("guardian") or {}).get("phone") or "",
                "since": oldest["due_date"] if oldest else "",
                "overdue": acc["overdue"],
                "balance": acc["balance"],
            }
        )
    rows.sort(key=lambda r: -r["overdue"])
    return {
        "title": f"Defaulters as of {on.isoformat()}",
        "columns": [
            ("prn", "PRN", "text"),
            ("student", "Student", "text"),
            ("class", "Class", "text"),
            ("phone", "Mobile", "text"),
            ("guardian_phone", "Guardian's mobile", "text"),
            ("since", "Overdue since", "date"),
            ("overdue", "Overdue", "money"),
            ("balance", "Balance", "money"),
        ],
        "rows": rows,
        "totals": {"Students": len(rows), "Overdue": sum(r["overdue"] for r in rows)},
    }


def scholarship_receivable(year_id: ObjectId) -> dict[str, Any]:
    db = get_db()
    rows = []
    stud = {s["_id"]: s for s in db.students.find({}, {"name": 1, "prn": 1})}
    for s in db.scholarships.find({"academic_year_id": year_id, "status": {"$in": ["sanctioned", "received"]}}):
        pending = s["sanctioned"] - s["received"]
        if pending <= 0:
            continue
        st = stud.get(s["student_id"], {})
        rows.append(
            {
                "scheme": s["scheme"],
                "prn": st.get("prn", ""),
                "student": st.get("name", ""),
                "reference": s.get("reference", ""),
                "sanctioned": s["sanctioned"],
                "received": s["received"],
                "pending": pending,
            }
        )
    rows.sort(key=lambda r: (r["scheme"], r["prn"]))
    by_scheme: dict[str, int] = {}
    for r in rows:
        by_scheme[r["scheme"]] = by_scheme.get(r["scheme"], 0) + r["pending"]
    return {
        "title": "Scholarships receivable",
        "columns": [
            ("scheme", "Scheme", "text"),
            ("prn", "PRN", "text"),
            ("student", "Student", "text"),
            ("reference", "Application ID", "text"),
            ("sanctioned", "Sanctioned", "money"),
            ("received", "Received", "money"),
            ("pending", "Still to come", "money"),
        ],
        "rows": rows,
        "totals": {"Still to come": sum(by_scheme.values()), **by_scheme},
    }


def build(name: str, params: dict[str, Any]) -> dict[str, Any]:
    year_id: ObjectId = params.get("year_id")  # type: ignore[assignment]  # set for the per-year reports
    if name == "day-book":
        return day_book(params.get("date_from"), params.get("date_to"))
    if name == "head-wise":
        return head_wise(params.get("date_from"), params.get("date_to"))
    if name == "mode-wise":
        return mode_wise(params.get("date_from"), params.get("date_to"))
    if name == "outstanding":
        return outstanding(year_id, params.get("programme_id"), params.get("year_of_study"))
    if name == "defaulters":
        return defaulters(year_id, params.get("programme_id"), params.get("year_of_study"), params.get("as_of"))
    if name == "scholarships":
        return scholarship_receivable(year_id)
    raise AppError(404, "No such report.")


def to_json(report: dict[str, Any]) -> dict[str, Any]:
    return {**report, "columns": [{"key": k, "header": h, "kind": t} for k, h, t in report["columns"]]}


def to_xlsx(report: dict[str, Any], subtitle: str) -> bytes:
    """One sheet: title, subtitle, header row, data, then totals. Money in rupees (2 decimals)."""
    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:31]
    ws.append([report["title"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([subtitle])
    ws.append([])
    ws.append([h for _, h, _ in report["columns"]])
    for cell in ws[4]:
        cell.font = Font(bold=True)
    money_cols = [i for i, (_, _, kind) in enumerate(report["columns"], start=1) if kind == "money"]
    for row in report["rows"]:
        ws.append([row[k] / 100 if kind == "money" else row[k] for k, _, kind in report["columns"]])
    for col in money_cols:
        for r in range(5, ws.max_row + 1):
            ws.cell(row=r, column=col).number_format = "#,##,##0.00"
    ws.append([])
    for label, value in report["totals"].items():
        is_money = isinstance(value, int) and label not in {"Students", "Cancelled receipts"}
        ws.append([label, value / 100 if is_money else value])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        if is_money:
            ws.cell(row=ws.max_row, column=2).number_format = "#,##,##0.00"
    for i, (_, header, _) in enumerate(report["columns"], start=1):
        width = max([len(header)] + [len(str(r[report["columns"][i - 1][0]])) for r in report["rows"][:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 40)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
