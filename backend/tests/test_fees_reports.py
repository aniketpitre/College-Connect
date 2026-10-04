import csv
import io
from datetime import date, timedelta

from openpyxl import load_workbook

from tests.fees_fixtures import RS, college, fees, new_client, structure_body  # noqa: F401

API = "/api/v1/fees"


def _setup(client, sign_in, fees, **kw):  # noqa: F811
    sign_in(client, ["accounts"])
    client.post(f"{API}/structures", json=structure_body(fees, **kw))
    client.post(
        f"{API}/demands/generate",
        json={"academic_year_id": fees["year"], "programme_id": fees["bca"], "year_of_study": 1, "dry_run": False},
    )


def _pay(client, fees, prn, amount, mode="cash", reference=""):  # noqa: F811
    body = {
        "student_id": fees["students"][prn],
        "academic_year_id": fees["year"],
        "amount": amount,
        "mode": mode,
        "reference": reference,
    }
    r = client.post(f"{API}/collect", json=body)
    assert r.status_code == 201, r.text
    return r.json()


def _csv(rows: list[list[str]], header=("prn", "paid", "old_receipt_no", "previous_dues", "note")) -> bytes:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(header)
    w.writerows(rows)
    return out.getvalue().encode()


def test_collection_reports(client, sign_in, fees):  # noqa: F811
    _setup(client, sign_in, fees)
    _pay(client, fees, "2026BCA001", 21_500 * RS)
    _pay(client, fees, "2026BCA002", 3_000 * RS, mode="upi", reference="UPI1")
    cancelled = _pay(client, fees, "2026BCA003", 1_000 * RS)
    req = client.post(f"{API}/receipts/{cancelled['id']}/cancel-request", json={"reason": "Wrong amount"}).json()
    principal = new_client(client)
    sign_in(principal, ["principal"])
    principal.post(f"/api/v1/approvals/{req['id']}/decide", json={"approve": True})
    today = client.get(f"{API}/today").json()["date"]

    book = client.get(f"{API}/reports/day-book", params={"date_from": today}).json()
    assert len(book["rows"]) == 3 and book["rows"][2]["status"] == "Cancelled" and book["rows"][2]["amount"] == 0
    assert book["totals"]["Collected"] == 24_500 * RS and book["totals"]["Cancelled receipts"] == 1

    heads = client.get(f"{API}/reports/head-wise", params={"date_from": today}).json()
    assert {r["code"]: r["amount"] for r in heads["rows"]} == {"TUITION": 23_000 * RS, "DEV": 1_500 * RS}
    modes = client.get(f"{API}/reports/mode-wise", params={"date_from": today}).json()
    assert {r["mode"]: (r["count"], r["amount"]) for r in modes["rows"]} == {
        "Cash": (1, 21_500 * RS),
        "UPI": (1, 3_000 * RS),
    }

    xlsx = client.get(f"{API}/reports/day-book", params={"date_from": today, "format": "xlsx"})
    assert xlsx.status_code == 200 and "spreadsheetml" in xlsx.headers["content-type"]
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws["A1"].value == "Day book" and ws["J5"].value == 21_500  # rupees, not paise
    assert client.get(f"{API}/reports/nope").status_code == 404


def test_outstanding_defaulters_and_scholarships(client, sign_in, fees):  # noqa: F811
    _setup(client, sign_in, fees, due_first=date.today() - timedelta(days=10))
    _pay(client, fees, "2026BCA001", 26_000 * RS)  # paid in full
    _pay(client, fees, "2026BCA002", 5_000 * RS)  # first installment (13,000) partly paid → overdue
    out = client.get(f"{API}/reports/outstanding").json()
    assert {r["prn"]: r["balance"] for r in out["rows"]} == {
        "2026BCA001": 0,
        "2026BCA002": 21_000 * RS,
        "2026BCA003": 26_000 * RS,
    }
    assert out["totals"]["Balance"] == 47_000 * RS

    late = client.get(f"{API}/reports/defaulters").json()
    assert [(r["prn"], r["overdue"]) for r in late["rows"]] == [("2026BCA003", 13_000 * RS), ("2026BCA002", 8_000 * RS)]
    early = client.get(
        f"{API}/reports/defaulters", params={"as_of": (date.today() - timedelta(days=20)).isoformat()}
    ).json()
    assert early["rows"] == []  # nothing was due yet

    sid = fees["students"]["2026BCA003"]
    s = client.post(
        f"{API}/scholarships",
        json={"student_id": sid, "academic_year_id": fees["year"], "scheme": "MahaDBT", "expected": 9_000 * RS},
    ).json()
    client.post(f"{API}/scholarships/{s['id']}/actions", json={"action": "sanction"})
    client.post(f"{API}/scholarships/{s['id']}/actions", json={"action": "receive", "amount": 4_000 * RS})
    rec = client.get(f"{API}/reports/scholarships").json()
    assert rec["rows"][0]["pending"] == 5_000 * RS and rec["totals"]["MahaDBT"] == 5_000 * RS


def test_opening_balances_import(client, sign_in, fees, db):  # noqa: F811
    _setup(client, sign_in, fees)
    assert client.get(f"{API}/opening/template.csv").text.startswith("prn,paid")
    bad = _csv(
        [
            ["2026BCA001", "13000", "1452", "", ""],
            ["2026BCA999", "100", "", "", ""],
            ["2026BCA002", "abc", "", "", ""],
            ["2026BCA003", "", "", "", ""],
            ["2026BCA001", "1", "", "", ""],
            ["2026BCA003", "99999", "", "", ""],
        ]
    )
    report = client.post(f"{API}/opening", files={"file": ("o.csv", bad, "text/csv")}, params={"commit": "true"}).json()
    assert report["committed"] is False and report["error_count"] == 5
    assert {(e["row"], e["field"]) for e in report["errors"]} == {
        (3, "prn"),
        (4, "paid"),
        (5, "paid"),
        (6, "prn"),
        (7, "prn"),
    }
    assert db.ledger_entries.count_documents({"type": {"$in": ["opening_paid", "opening_due"]}}) == 0  # nothing saved

    over = _csv([["2026BCA003", "99999", "", "", ""]])
    too_much = client.post(f"{API}/opening", files={"file": ("o.csv", over, "text/csv")}).json()
    assert too_much["errors"][0]["message"].startswith("More than the")

    good = _csv([["2026BCA001", "13,000", "1452", "", "Paid in June"], ["2026BCA002", "", "", "2500", "From 2025-26"]])
    preview = client.post(f"{API}/opening", files={"file": ("o.csv", good, "text/csv")}).json()
    assert preview["committed"] is False and preview["paid"] == 13_000 * RS and preview["dues"] == 2_500 * RS
    done = client.post(f"{API}/opening", files={"file": ("o.csv", good, "text/csv")}, params={"commit": "true"}).json()
    assert done["committed"] is True
    rohan = client.get(f"{API}/students/{fees['students']['2026BCA001']}").json()
    assert rohan["balance"] == 13_000 * RS and rohan["entries"][-1]["reason"] == "Old receipt 1452 · Paid in June"
    neha = client.get(f"{API}/students/{fees['students']['2026BCA002']}").json()
    assert neha["balance"] == 28_500 * RS and "PREV" in [h["code"] for h in neha["by_head"]]
    again = client.post(f"{API}/opening", files={"file": ("o.csv", good, "text/csv")}, params={"commit": "true"}).json()
    assert again["committed"] is False and "already imported" in again["errors"][0]["message"]


def test_reports_permissions(client, sign_in, fees):  # noqa: F811
    office = new_client(client)
    sign_in(office, ["office"])
    assert office.get(f"{API}/reports/outstanding").status_code == 200
    assert office.post(f"{API}/opening", files={"file": ("o.csv", b"prn,paid\n", "text/csv")}).status_code == 403
    faculty = new_client(client)
    sign_in(faculty, ["faculty"])
    assert faculty.get(f"{API}/reports/day-book").status_code == 403
